import csv
from datetime import timedelta
from io import BytesIO, StringIO

from django.contrib.auth import get_user_model
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework import status
from rest_framework.test import APITestCase

from chatbot.models import Chatbot, ChatbotCapacity, ChatbotUser
from chatbot.utils.choices import ChatbotRoleTypes
from lead_capture.models import Lead, LeadCaptureConfig, LeadNote
from subscription.choices import PlanFeature
from workspace.models import Workspace, WorkspaceRole, WorkspaceUser

User = get_user_model()


class LeadCaptureClientAPITests(APITestCase):
    def setUp(self):
        self.user = User.objects.create_user(
            email="lead-api@example.com",
            password="StrongPass123!",
        )
        self.workspace = Workspace.objects.create(
            name="Lead API Workspace",
            slug="lead-api-workspace",
            owner=self.user,
        )
        WorkspaceUser.objects.create(
            workspace=self.workspace,
            user=self.user,
            role=WorkspaceRole.ADMIN,
        )
        self.chatbot = Chatbot.objects.create(
            workspace=self.workspace,
            chatbot_name="Lead API Bot",
            slug="lead-api-bot",
            created_by=self.user,
        )
        self.chatbot_user = ChatbotUser.objects.create(
            chatbot=self.chatbot,
            user=self.user,
            role=ChatbotRoleTypes.ADMIN,
        )
        self.capacity = ChatbotCapacity.objects.create(
            chatbot=self.chatbot,
            active_features=[PlanFeature.LEAD_CAPTURE],
        )
        self.client.force_authenticate(self.user)

    def url(self, name, *, lead=None, note=None):
        query = f"?chatbot_slug={self.chatbot.slug}"
        if lead is not None:
            query += f"&lead_id={lead.id}"
        if note is not None:
            query += f"&note_id={note.id}"
        return f"{reverse(name)}{query}"

    def test_configuration_can_be_created_fetched_and_updated(self):
        response = self.client.post(
            self.url("lead-config-create"),
            {
                "is_enabled": True,
                "collectable_fields": [
                    {
                        "label": "Organization",
                        "value": "organization",
                        "mode": "required",
                        "type": "text",
                    },
                    {
                        "label": "Team Size",
                        "value": "team_size",
                        "mode": "optional",
                        "type": "text",
                    },
                ],
            },
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        config = LeadCaptureConfig.objects.get(chatbot=self.chatbot)
        self.assertTrue(config.is_enabled)
        self.assertEqual(len(config.collectable_fields), 2)

        response = self.client.get(self.url("lead-config"))
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data["data"]["chatbot_id"], str(self.chatbot.id))

        response = self.client.patch(
            self.url("lead-config-update"),
            {},
            format="json",
        )
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        config.refresh_from_db()

    def test_feature_is_required(self):
        self.capacity.active_features = []
        self.capacity.save(update_fields=["active_features", "updated_at"])

        response = self.client.get(self.url("lead-list"))

        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)

    def test_leads_are_paginated(self):
        for number in range(3):
            Lead.objects.create(
                chatbot=self.chatbot,
                collected_fields={
                    "name": f"Lead {number}",
                    "email": f"lead-{number}@example.com",
                },
            )

        response = self.client.get(f'{self.url("lead-list")}&page_size=2')

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data["meta"]["count"], 3)
        self.assertEqual(len(response.data["data"]), 2)
        self.assertIn("notes_count", response.data["data"][0])

    def test_lead_stats_returns_totals_and_breakdowns_for_chatbot(self):
        Lead.objects.bulk_create(
            [
                Lead(
                    chatbot=self.chatbot,
                    status="new",
                    lead_score=75,
                    source="widget",
                ),
                Lead(
                    chatbot=self.chatbot,
                    status="qualified",
                    lead_score=76,
                    source="widget",
                ),
                Lead(
                    chatbot=self.chatbot,
                    status="converted",
                    lead_score=100,
                    source="campaign",
                ),
                Lead(
                    chatbot=self.chatbot,
                    status="new",
                    lead_score=None,
                    source="",
                ),
            ]
        )
        other_chatbot = Chatbot.objects.create(
            workspace=self.workspace,
            chatbot_name="Other Lead Bot",
            created_by=self.user,
        )
        Lead.objects.create(
            chatbot=other_chatbot,
            lead_score=99,
            source="other",
        )

        response = self.client.get(self.url("lead-stats"))

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        data = response.data["data"]
        self.assertEqual(data["total_leads"], 4)
        self.assertEqual(data["hot_leads"], 2)
        self.assertEqual(data["hot_lead_percentage"], 50.0)
        self.assertEqual(data["scored_leads"], 3)
        self.assertEqual(data["unscored_leads"], 1)
        self.assertEqual(data["average_lead_score"], 83.67)
        self.assertEqual(
            data["leads_by_channel"],
            [
                {"channel": "widget", "count": 2, "percentage": 50.0},
                {"channel": "unknown", "count": 1, "percentage": 25.0},
                {"channel": "campaign", "count": 1, "percentage": 25.0},
            ],
        )
        self.assertEqual(
            {
                item["status"]: item["count"]
                for item in data["leads_by_status"]
            },
            {
                "new": 2,
                "qualified": 1,
                "contacted": 0,
                "converted": 1,
                "disqualified": 0,
            },
        )

    def test_lead_stats_returns_defined_empty_values(self):
        response = self.client.get(self.url("lead-stats"))

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        data = response.data["data"]
        self.assertEqual(data["total_leads"], 0)
        self.assertEqual(data["hot_leads"], 0)
        self.assertEqual(data["hot_lead_percentage"], 0.0)
        self.assertEqual(data["scored_leads"], 0)
        self.assertEqual(data["unscored_leads"], 0)
        self.assertIsNone(data["average_lead_score"])
        self.assertEqual(data["leads_by_channel"], [])
        self.assertTrue(
            all(item["count"] == 0 for item in data["leads_by_status"])
        )

    def test_all_leads_can_be_exported_as_csv_without_a_date_range(self):
        Lead.objects.create(
            chatbot=self.chatbot,
            collected_fields={
                "name": "CSV Lead",
                "email": "csv@example.com",
            },
            initial_ip_address="192.0.2.10",
            detected_country_code="US",
            detected_city="New York",
            source="widget",
        )

        response = self.client.get(
            f'{self.url("export-lead-data")}&file_format=csv'
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response["Content-Type"], "text/csv; charset=utf-8")
        self.assertIn("attachment;", response["Content-Disposition"])
        self.assertEqual(response["X-Lead-Export-Limit"], "1000")
        self.assertEqual(response["X-Lead-Export-Count"], "1")
        rows = list(csv.reader(StringIO(response.content.decode("utf-8-sig"))))
        self.assertEqual(len(rows), 2)
        self.assertIn("name", rows[0])
        self.assertIn("email", rows[0])
        self.assertEqual(rows[1][rows[0].index("name")], "CSV Lead")
        self.assertEqual(
            rows[1][rows[0].index("email")],
            "csv@example.com",
        )

    def test_lead_export_filters_by_inclusive_date_range(self):
        old_lead = Lead.objects.create(
            chatbot=self.chatbot,
            collected_fields={"name": "Old Lead"},
        )
        current_lead = Lead.objects.create(
            chatbot=self.chatbot,
            collected_fields={"name": "Current Lead"},
        )
        Lead.objects.filter(pk=old_lead.pk).update(
            created_at=timezone.now() - timedelta(days=5),
        )
        today = timezone.localdate().isoformat()

        response = self.client.get(
            f'{self.url("export-lead-data")}&file_format=csv'
            f"&start_date={today}&end_date={today}"
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        rows = list(csv.reader(StringIO(response.content.decode("utf-8-sig"))))
        self.assertEqual(response["X-Lead-Export-Count"], "1")
        self.assertEqual(
            rows[1][rows[0].index("name")],
            current_lead.collected_fields["name"],
        )

    def test_leads_can_be_exported_as_excel(self):
        Lead.objects.create(
            chatbot=self.chatbot,
            collected_fields={"name": "Excel Lead"},
        )

        response = self.client.get(
            f'{self.url("export-lead-data")}&file_format=excel'
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(
            response["Content-Type"],
            (
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            ),
        )
        self.assertIn(".xlsx", response["Content-Disposition"])
        workbook = load_workbook(BytesIO(response.content), read_only=True)
        rows = list(workbook["Leads"].iter_rows(values_only=True))
        self.assertIn("name", rows[0])
        self.assertEqual(rows[1][rows[0].index("name")], "Excel Lead")

    def test_lead_export_has_a_hard_limit_of_1000(self):
        Lead.objects.bulk_create(
            [
                Lead(
                    chatbot=self.chatbot,
                    collected_fields={"name": f"Lead {index}"},
                )
                for index in range(1001)
            ]
        )

        response = self.client.get(
            f'{self.url("export-lead-data")}&file_format=csv'
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        rows = list(csv.reader(StringIO(response.content.decode("utf-8-sig"))))
        self.assertEqual(response["X-Lead-Export-Limit"], "1000")
        self.assertEqual(response["X-Lead-Export-Count"], "1000")
        self.assertEqual(len(rows), 1001)

    def test_lead_export_validates_format_and_date_order(self):
        response = self.client.get(
            f'{self.url("export-lead-data")}&file_format=pdf'
        )
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)

        response = self.client.get(
            f'{self.url("export-lead-data")}&file_format=csv'
            "&start_date=2026-09-25&end_date=2026-09-24"
        )
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("end_date", response.data["errors"])

    def test_lead_information_can_be_updated(self):
        LeadCaptureConfig.objects.create(
            chatbot=self.chatbot,
            collectable_fields=[
                {
                    "label": "Name",
                    "value": "name",
                    "mode": "required",
                    "type": "text",
                },
                {
                    "label": "Email",
                    "value": "email",
                    "mode": "required",
                    "type": "email",
                },
                {
                    "label": "Organization",
                    "value": "organization",
                    "mode": "optional",
                    "type": "text",
                },
            ],
        )
        lead = Lead.objects.create(
            chatbot=self.chatbot,
            collected_fields={
                "name": "Original Name",
                "email": "original@example.com",
            },
        )

        response = self.client.patch(
            self.url("lead-update", lead=lead),
            {
                "collected_fields": {
                    "name": "Updated Name",
                    "email": "UPDATED@EXAMPLE.COM",
                    "organization": "Argon",
                },
                "status": "contacted",
                "lead_score": 75,
            },
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        lead.refresh_from_db()
        self.assertEqual(lead.collected_fields["name"], "Updated Name")
        self.assertEqual(lead.collected_fields["email"], "updated@example.com")
        self.assertEqual(lead.status, "contacted")
        self.assertEqual(lead.lead_score, 75)

    def test_notes_can_be_created_fetched_updated_listed_and_deleted(self):
        lead = Lead.objects.create(
            chatbot=self.chatbot,
            collected_fields={"name": "Noted Lead"},
        )

        response = self.client.post(
            self.url("lead-note-create", lead=lead),
            {"content": "Follow up tomorrow."},
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        note = LeadNote.objects.get(lead=lead)
        self.assertEqual(note.author, self.chatbot_user)
        self.assertEqual(note.content, "Follow up tomorrow.")

        response = self.client.get(
            self.url("lead-note-detail", lead=lead, note=note)
        )
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data["data"]["id"], str(note.id))

        response = self.client.patch(
            self.url("lead-note-update", lead=lead, note=note),
            {"content": "Follow up next week."},
            format="json",
        )
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        note.refresh_from_db()
        self.assertEqual(note.content, "Follow up next week.")

        response = self.client.get(
            self.url("lead-note-list", lead=lead)
        )
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data["meta"]["count"], 1)
        self.assertEqual(
            response.data["data"][0]["author"]["email"],
            self.user.email,
        )

        response = self.client.delete(
            self.url("lead-note-delete", lead=lead, note=note)
        )
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data["data"]["id"], str(note.id))
        self.assertFalse(LeadNote.objects.filter(pk=note.id).exists())

    def test_note_detail_is_scoped_to_its_lead(self):
        lead = Lead.objects.create(
            chatbot=self.chatbot,
            collected_fields={"name": "First Lead"},
        )
        other_lead = Lead.objects.create(
            chatbot=self.chatbot,
            collected_fields={"name": "Second Lead"},
        )
        note = LeadNote.objects.create(
            lead=lead,
            author=self.chatbot_user,
            content="Private to the first lead.",
        )

        response = self.client.get(
            self.url("lead-note-detail", lead=other_lead, note=note)
        )

        self.assertEqual(response.status_code, status.HTTP_404_NOT_FOUND)

    def test_task_endpoints_only_accept_their_own_http_methods(self):
        lead = Lead.objects.create(
            chatbot=self.chatbot,
            collected_fields={"name": "Method Lead"},
        )

        self.assertEqual(
            self.client.post(self.url("lead-config")).status_code,
            status.HTTP_405_METHOD_NOT_ALLOWED,
        )
        self.assertEqual(
            self.client.patch(
                self.url("lead-detail", lead=lead),
                {},
                format="json",
            ).status_code,
            status.HTTP_405_METHOD_NOT_ALLOWED,
        )
        self.assertEqual(
            self.client.post(
                self.url("lead-note-list", lead=lead),
                {"content": "Wrong endpoint."},
                format="json",
            ).status_code,
            status.HTTP_405_METHOD_NOT_ALLOWED,
        )
